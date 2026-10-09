"""AUTO-GENERATED from 06_Enforcement_Letters_and_Summary.ipynb by build_steps.py - edit the notebook and regenerate, not this file."""
# ----------------------------------------------------------------------
# Connects Google Drive, finds your project folder and loads the shared helper module `map_common`.
# ----------------------------------------------------------------------
import os
import map_common as mc          # run_step.py has already put the project folder on the import path
mc.ensure_dirs()
print("Project folder:", mc.PROJECT_DIR)

# ----------------------------------------------------------------------
# Starts MySQL (restoring the Drive backup in a new Colab session) and connects.
# ----------------------------------------------------------------------
import os, re, zipfile, textwrap
import pandas as pd
from sqlalchemy import text
engine = mc.connect_db()

# ----------------------------------------------------------------------
# Settings: which date(s) to process, the strike at which a formal letter replaces warnings, the letter format (text, PDF or both) and the wording place
# ----------------------------------------------------------------------
RUN_DATES = _ov("RUN_DATES", None)            # None = latest date with violations | ["2026-01-05"] | "ALL" (every date - creates many files)
LETTER_AT = _ov("LETTER_AT", 3)               # strike number at which a formal Violation Letter replaces warnings
LETTER_FORMAT = _ov("LETTER_FORMAT", "txt")       # "txt" | "pdf" | "both"

MANUFACTURER_NAME = _ov("MANUFACTURER_NAME", "[Manufacturer Name]")                      # <-- replace before sending
CONTACT_LINE = _ov("CONTACT_LINE", "[Channel Compliance Team | e-mail | phone]")     # <-- replace before sending
RESPONSE_DAYS = _ov("RESPONSE_DAYS", 7)           # days the seller has to correct the listings
assert LETTER_FORMAT in ("txt", "pdf", "both")
if LETTER_FORMAT in ("pdf", "both"):
    mc.ensure_packages({"reportlab": "reportlab"})
print("Letter threshold: strike", LETTER_AT, "| format:", LETTER_FORMAT)

# ----------------------------------------------------------------------
# Reads every violating observation from `PRICE_MONITORING_TABLE` and uses the recognised (homologated) seller name, falling back to the raw name when n
# ----------------------------------------------------------------------
q = """SELECT SKU, Violation_date, PL, Category, Sub_category, Region, Marketplace, seller_name, homologated_name,
              MAP_Price, LPP, Advertised_price, currency_local, Advertised_price_base, violation_reason
       FROM PRICE_MONITORING_TABLE WHERE is_violation = 1"""
with engine.connect() as conn:
    violations = pd.read_sql(text(q), conn)
if violations.empty:
    raise RuntimeError("There are no violations in PRICE_MONITORING_TABLE (run notebook 05, or nothing was below the floor).")
for c in ["MAP_Price", "LPP", "Advertised_price", "Advertised_price_base"]:
    violations[c] = pd.to_numeric(violations[c], errors="coerce")
violations["Violation_date"] = pd.to_datetime(violations["Violation_date"]).dt.normalize()
violations["Seller"] = violations["homologated_name"].fillna(violations["seller_name"])
print(f"Violating observations: {len(violations):,} | dates {violations['Violation_date'].min().date()} to {violations['Violation_date'].max().date()}")
print(f"Recognised sellers: {violations['Seller'].nunique()} | seller + marketplace pairs: {violations.groupby(['Seller', 'Marketplace']).ngroups}")

# ----------------------------------------------------------------------
# Defines the strike logic: one row per seller + marketplace + day, the cumulative strike count up to the run date, and the action for that strike. Test
# ----------------------------------------------------------------------
def action_for(strike, letter_at=None):
    letter_at = LETTER_AT if letter_at is None else letter_at
    return "Violation Letter" if strike >= letter_at else f"Warning {strike}"

def build_daily(v):
    """One row per (Seller, Marketplace, day) listing the SKUs involved."""
    return (v.groupby(["Seller", "Marketplace", "Violation_date"])["SKU"]
              .agg(lambda s: sorted(set(s))).rename("SKUs").reset_index())

def actions_for_date(daily, run_date, letter_at=None):
    run_date = pd.Timestamp(run_date).normalize()
    history = daily[daily["Violation_date"] <= run_date]
    strikes = history.groupby(["Seller", "Marketplace"]).size().rename("Strike").reset_index()
    today = daily[daily["Violation_date"] == run_date].merge(strikes, on=["Seller", "Marketplace"])
    today["Action"] = today["Strike"].apply(lambda s: action_for(s, letter_at))
    return today.sort_values(["Seller", "Marketplace"]).reset_index(drop=True)

# --- test with invented data: A on Mkt1 violates on days 1,2,3,4 (several rows per day); A on Mkt2 only day 4; B only day 4
d = pd.Timestamp
test = pd.DataFrame([("A", "Mkt1", d("2026-01-01"), "X1"), ("A", "Mkt1", d("2026-01-01"), "X2"), ("A", "Mkt1", d("2026-01-02"), "X1"),
                     ("A", "Mkt1", d("2026-01-03"), "X1"), ("A", "Mkt1", d("2026-01-04"), "X1"),
                     ("A", "Mkt2", d("2026-01-04"), "X1"), ("B", "Mkt1", d("2026-01-04"), "X9")], columns=["Seller", "Marketplace", "Violation_date", "SKU"])
td = build_daily(test)
got = {day: actions_for_date(td, day, 3) for day in ["2026-01-01", "2026-01-02", "2026-01-03", "2026-01-04"]}
assert got["2026-01-01"].loc[0, "Action"] == "Warning 1" and got["2026-01-01"].loc[0, "Strike"] == 1   # 2 rows same day = 1 strike
assert got["2026-01-02"].loc[0, "Action"] == "Warning 2"
assert got["2026-01-03"].loc[0, "Action"] == "Violation Letter" and got["2026-01-03"].loc[0, "Strike"] == 3
day4 = got["2026-01-04"].set_index(["Seller", "Marketplace"])
assert day4.loc[("A", "Mkt1"), "Strike"] == 4 and day4.loc[("A", "Mkt2"), "Action"] == "Warning 1" and day4.loc[("B", "Mkt1"), "Action"] == "Warning 1"
print("Strike logic tests passed")

daily = build_daily(violations)
if RUN_DATES is None:
    dates_to_run = [violations["Violation_date"].max()]
elif RUN_DATES == "ALL":
    dates_to_run = sorted(violations["Violation_date"].unique())
else:
    dates_to_run = [pd.Timestamp(x) for x in RUN_DATES]
print("Run date(s):", [pd.Timestamp(x).date().isoformat() for x in dates_to_run])

# ----------------------------------------------------------------------
# Defines `build_letter()`, which writes the letter text: header, seller, marketplace, strike and action, a tone that depends on the action (reminder fo
# ----------------------------------------------------------------------
def money(x):
    return "n/a" if pd.isna(x) else f"{x:,.2f}"

def safe_name(s):
    return re.sub(r"[^\w\-]+", "_", str(s)).strip("_")

def build_letter(seller, marketplace, run_date, strike, action, day_rows):
    d = pd.Timestamp(run_date).strftime("%d %B %Y")
    lines = []
    for r in day_rows.sort_values("SKU").itertuples():
        lines.append(
            f"  - SKU {r.SKU} ({r.PL}) | Region: {r.Region}\n"
            f"      Advertised: {r.currency_local} {money(r.Advertised_price)}  (= INR {money(r.Advertised_price_base)})\n"
            f"      MAP: INR {money(r.MAP_Price)} | LPP: INR {money(r.LPP)}\n"
            f"      Reason: {r.violation_reason}")
    detail = "\n".join(lines)
    if action == "Violation Letter":
        subject = "FORMAL NOTICE OF PRICING POLICY VIOLATION"
        body = (f"This is a formal notice that {seller} has advertised {MANUFACTURER_NAME} products on {marketplace} below the "
                f"permitted price on {strike} separate days to date, and has previously been reminded of the pricing policy. "
                f"This is violation strike number {strike}.\n\n"
                f"Please remove or correct the listings below within {RESPONSE_DAYS} days of this notice and confirm in writing. "
                f"Further non-compliance may result in the suspension of supply, promotional eligibility or authorisation to sell "
                f"{MANUFACTURER_NAME} products.")
    else:
        subject = f"PRICING POLICY REMINDER - {action.upper()}"
        body = (f"Our monitoring found that {seller} advertised {MANUFACTURER_NAME} products on {marketplace} below the permitted price. "
                f"This is {action} (strike number {strike}).\n\n"
                f"Please review and correct the listings below within {RESPONSE_DAYS} days. "
                f"Repeated violations will lead to a formal Violation Letter.")
    return (f"{MANUFACTURER_NAME}\n{'=' * 70}\n{subject}\n{'=' * 70}\n\n"
            f"Date:        {d}\nTo:          {seller}\nMarketplace: {marketplace}\nStrike:      {strike}\nAction:      {action}\n\n"
            f"Dear {seller},\n\n{body}\n\nListings found on {d}:\n{detail}\n\n"
            f"Pricing definitions: MAP is the minimum advertised price. LPP is the lowest possible price and may never be undercut, "
            f"even during a promotion.\n\nIf you believe this is an error, please contact {CONTACT_LINE}.\n\n"
            f"Sincerely,\n{MANUFACTURER_NAME} Channel Compliance\n")
print("Letter template ready.")

# ----------------------------------------------------------------------
# Defines `write_pdf()` (puts a letter on A4 pages in a fixed-width font, wrapping long lines) and `run_for_date()`: for one date it finds who gets whic
# ----------------------------------------------------------------------
def write_pdf(path, content):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(path, pagesize=A4)
    w, h = A4; y = h - 50
    c.setFont("Courier", 9)
    for raw in content.splitlines():
        raw = raw.encode("latin-1", "replace").decode("latin-1")
        indent = " " * (len(raw) - len(raw.lstrip()))
        for part in (textwrap.wrap(raw, 88, subsequent_indent=indent + "  ", drop_whitespace=False) or [""]):
            if y < 50:
                c.showPage(); c.setFont("Courier", 9); y = h - 50
            c.drawString(50, y, part.rstrip()); y -= 12
    c.save()

def run_for_date(run_date):
    run_date = pd.Timestamp(run_date).normalize()
    date_str = run_date.strftime("%Y-%m-%d")
    todays = actions_for_date(daily, run_date)
    day_viol = violations[violations["Violation_date"] == run_date]
    out, written = [], []
    for r in todays.itertuples():
        rows = day_viol[(day_viol["Seller"] == r.Seller) & (day_viol["Marketplace"] == r.Marketplace)]
        folder = os.path.join(mc.LETTER_DIR, r.Action.replace(" ", "_")); os.makedirs(folder, exist_ok=True)
        base = os.path.join(folder, f"{safe_name(r.Seller)}_{safe_name(r.Marketplace)}_{date_str}")
        content = build_letter(r.Seller, r.Marketplace, run_date, r.Strike, r.Action, rows)
        paths = []
        if LETTER_FORMAT in ("txt", "both"):
            with open(base + ".txt", "w", encoding="utf-8") as f: f.write(content)
            paths.append(base + ".txt")
        if LETTER_FORMAT in ("pdf", "both"):
            write_pdf(base + ".pdf", content); paths.append(base + ".pdf")
        written += paths
        out.append({"Date": run_date, "Seller (homologated_name)": r.Seller, "Marketplace": r.Marketplace,
                    "Strike Number": r.Strike, "Action": r.Action, "SKUs Involved": ", ".join(r.SKUs),
                    "File Path": ", ".join(os.path.relpath(p, mc.PROJECT_DIR) for p in paths)})
    summary = pd.DataFrame(out, columns=["Date", "Seller (homologated_name)", "Marketplace", "Strike Number", "Action", "SKUs Involved", "File Path"])

    xlsx = os.path.join(mc.SUMMARY_DIR, f"Summary_{date_str}.xlsx")
    with pd.ExcelWriter(xlsx, engine="openpyxl", datetime_format="yyyy-mm-dd") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        ws = xw.sheets["Summary"]
        from openpyxl.styles import Font, PatternFill, Alignment
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for col, wd in zip("ABCDEFG", [12, 32, 16, 14, 18, 50, 70]):
            ws.column_dimensions[col].width = wd
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    zpath = os.path.join(mc.SUMMARY_DIR, f"Letters_{date_str}.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in written:
            z.write(p, os.path.relpath(p, mc.LETTER_DIR))
    print(f"{date_str}: {len(summary)} letter(s)/warning(s) -> {os.path.relpath(xlsx, mc.PROJECT_DIR)} and {os.path.basename(zpath)}")
    return summary
print("Writers ready.")

# ----------------------------------------------------------------------
# Runs the enforcement for each selected date.
# ----------------------------------------------------------------------
summaries = {pd.Timestamp(x).date().isoformat(): run_for_date(x) for x in dates_to_run}

# ----------------------------------------------------------------------
# Adds every letter just produced to the approval queue (`ENFORCEMENT_QUEUE`) with status `Pending`, together with the seller's e-mail address from `SEL
# ----------------------------------------------------------------------
QUEUE_DDL = [
"""CREATE TABLE IF NOT EXISTS SELLER_CONTACT (
    homologated_name VARCHAR(255) COLLATE utf8mb4_bin NOT NULL,
    marketplace      VARCHAR(150) COLLATE utf8mb4_bin NOT NULL DEFAULT '*',
    email            VARCHAR(255),
    contact_person   VARCHAR(255),
    PRIMARY KEY (homologated_name, marketplace)) ENGINE=InnoDB""",
"""CREATE TABLE IF NOT EXISTS ENFORCEMENT_QUEUE (
    queue_id         INT AUTO_INCREMENT PRIMARY KEY,
    run_date         DATE NOT NULL,
    seller           VARCHAR(255) COLLATE utf8mb4_bin NOT NULL,
    marketplace      VARCHAR(150) COLLATE utf8mb4_bin NOT NULL,
    strike           INT,
    action           VARCHAR(30),
    skus             TEXT,
    letter_path      VARCHAR(1000),
    recipient_email  VARCHAR(255),
    status           VARCHAR(20) NOT NULL DEFAULT 'Pending',
    comment          VARCHAR(500),
    approved_by      VARCHAR(100),
    approved_at      DATETIME,
    sent_at          DATETIME,
    gmail_message_id VARCHAR(100),
    error            VARCHAR(500),
    created_at       DATETIME DEFAULT CURRENT_TIMESTAMP,
    UNIQUE KEY UQ_QUEUE (run_date, seller, marketplace)) ENGINE=InnoDB""",
]
with engine.begin() as conn:
    for stmt in QUEUE_DDL:
        conn.execute(text(stmt))
    contacts = pd.read_sql(text("SELECT homologated_name, marketplace, email FROM SELLER_CONTACT"), conn)

def find_email(seller, marketplace):
    hit = contacts[(contacts["homologated_name"] == seller) & (contacts["marketplace"] == marketplace)]
    if hit.empty:
        hit = contacts[(contacts["homologated_name"] == seller) & (contacts["marketplace"] == "*")]
    return hit["email"].iloc[0] if len(hit) else None

UPSERT = text("""
    INSERT INTO ENFORCEMENT_QUEUE (run_date, seller, marketplace, strike, action, skus, letter_path, recipient_email, status)
    VALUES (:run_date, :seller, :marketplace, :strike, :action, :skus, :letter_path, :email, 'Pending')
    ON DUPLICATE KEY UPDATE
        status = IF(status = 'Sent', status, IF(action <> VALUES(action), 'Pending', status)),
        strike = VALUES(strike), action = VALUES(action), skus = VALUES(skus), letter_path = VALUES(letter_path),
        recipient_email = COALESCE(recipient_email, VALUES(recipient_email))""")
queued = 0
with engine.begin() as conn:
    for date_str, s_ in summaries.items():
        for rec in s_.to_dict("records"):
            seller, mkt = rec["Seller (homologated_name)"], rec["Marketplace"]
            paths = rec["File Path"].split(", ")
            main_path = next((p for p in paths if p.endswith(".txt")), paths[0])
            conn.execute(UPSERT, {"run_date": date_str, "seller": seller, "marketplace": mkt, "strike": int(rec["Strike Number"]),
                                  "action": rec["Action"], "skus": rec["SKUs Involved"], "letter_path": main_path,
                                  "email": find_email(seller, mkt)})
            queued += 1
    counts = pd.read_sql(text("SELECT status, COUNT(*) AS letters FROM ENFORCEMENT_QUEUE GROUP BY status"), conn)
print(f"Queued {queued} letter(s) for review. Nothing has been sent.")
print(counts.to_string(index=False))

# ----------------------------------------------------------------------
# Verifies the last run: counts by action and strike, number of distinct sellers, and that every letter file listed in the summary really exists on Driv
# ----------------------------------------------------------------------
last = sorted(summaries)[-1]; s = summaries[last]
print(f"Verification for {last}")
print("\nRows by action:\n" + s["Action"].value_counts().to_string())
print("\nStrike spread:\n" + s.groupby(["Action", "Strike Number"]).size().to_string())
missing = [p for cell in s["File Path"] for p in cell.split(", ") if not os.path.exists(os.path.join(mc.PROJECT_DIR, p))]
print(f"\nSummary rows: {len(s)} | distinct sellers: {s['Seller (homologated_name)'].nunique()} | letter files missing on disk: {len(missing)}")
assert not missing
expected_pairs = daily[daily["Violation_date"] == pd.Timestamp(last)].shape[0]
assert len(s) == expected_pairs, f"{len(s)} summary rows but {expected_pairs} seller+marketplace pairs violated that day"
s.head(10)

# ----------------------------------------------------------------------
# Shows the text of one Violation Letter (or, if none today, one warning) so you can read exactly what a seller would receive.
# ----------------------------------------------------------------------
pick = s[s["Action"] == "Violation Letter"].head(1)
if pick.empty:
    pick = s.head(1)
first_path = pick["File Path"].iloc[0].split(", ")[0]
if first_path.endswith(".txt"):
    print(open(os.path.join(mc.PROJECT_DIR, first_path), encoding="utf-8").read())
else:
    print("Letter written as PDF:", first_path)
